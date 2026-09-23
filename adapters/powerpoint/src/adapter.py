#!/usr/bin/env python3
"""Offline PowerPoint adapter over Open XML (python-pptx).

Route: exact .pptx file in, mutate through a closed op schema, save,
reopen, and verify by readback. Macros/VBA are never executed.
"""
import hashlib
import math
import os
import re
import shutil
import subprocess
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "_shared"))
from adapter_protocol import response, serve  # noqa: E402

ADAPTER_ID = "comptrol.powerpoint"
READ_EXTS = {".pptx", ".pptm"}
WRITE_EXTS = {".pptx"}
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".gif", ".bmp"}
EXPORT_FORMATS = {"pdf", "pptx"}
BACKUP_SUFFIX = ".bak"
MAX_OPS = 50
MACRO_KEYS = {"macro", "macros", "vba", "vba_macro", "run_macro", "execute_macro", "vbproject", "oleobject"}


def has_pptx():
    try:
        import pptx  # noqa: F401
        return True
    except ImportError:
        return False


def soffice_bin():
    override = os.environ.get("COMPTROL_SOFFICE_BIN", "").strip()
    if override:
        candidate = Path(override)
        if candidate.is_file():
            return str(candidate)
        return None
    for name in ("soffice", "libreoffice"):
        found = shutil.which(name)
        if found:
            return found
    return None


def scope_root():
    raw = os.environ.get("COMPTROL_PRESENTATIONS_ROOT", "").strip()
    if not raw:
        return None
    return Path(raw).resolve()


def scoped_resolve(raw, allowed_exts=None, must_exist=True):
    if not isinstance(raw, str) or not raw.strip():
        raise ValueError("path is required")
    candidate = Path(raw.strip()).resolve()
    root = scope_root()
    if root is not None:
        try:
            candidate.relative_to(root)
        except ValueError:
            raise ValueError("path_out_of_scope: path is outside COMPTROL_PRESENTATIONS_ROOT")
    if allowed_exts is not None and candidate.suffix.lower() not in allowed_exts:
        raise ValueError("unsupported extension for %s (got %s)" % (sorted(allowed_exts), candidate.suffix))
    if must_exist and not candidate.is_file():
        raise ValueError("file not found: %s" % candidate)
    return candidate


def sha256_of(path):
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        for chunk in iter(lambda: handle.read(65536), b""):
            digest.update(chunk)
    return digest.hexdigest()


def reject_macro_requests(payload):
    if not isinstance(payload, dict):
        return
    for key in payload.keys():
        if isinstance(key, str) and key.strip().lower() in MACRO_KEYS:
            raise ValueError("macro_execution_refused: VBA/macros are never executed")
    ops = payload.get("ops")
    if isinstance(ops, list):
        for op in ops:
            if not isinstance(op, dict):
                continue
            for key in op.keys():
                if isinstance(key, str) and key.strip().lower() in MACRO_KEYS:
                    raise ValueError("macro_execution_refused: VBA/macros are never executed")
            blob = " ".join(str(value).lower() for value in op.values() if isinstance(value, str))
            if "vba" in blob and "macro" in blob and "run" in blob:
                raise ValueError("macro_execution_refused: VBA/macros are never executed")


def shape_texts(shape):
    texts = []
    try:
        if shape.has_text_frame:
            for paragraph in shape.text_frame.paragraphs:
                texts.append("".join(run.text for run in paragraph.runs))
        if shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    texts.append(cell.text)
    except Exception:
        pass
    return texts


def slide_text_content(slide):
    texts = []
    for shape in slide.shapes:
        texts.extend(shape_texts(shape))
    return texts


def notes_text(slide):
    try:
        return slide.notes_slide.notes_text_frame.text
    except Exception:
        return ""


def summarize(prs, path):
    slides = []
    for index, slide in enumerate(prs.slides, start=1):
        slides.append(
            {
                "index": index,
                "shape_count": len(slide.shapes),
                "texts": slide_text_content(slide),
                "notes": notes_text(slide),
            }
        )
    size = path.stat().st_size if path.is_file() else 0
    return {
        "path": str(path),
        "slide_count": len(prs.slides),
        "slides": slides,
        "file_size": size,
        "sha256": sha256_of(path) if path.is_file() and size > 0 else "",
    }


def parse_slide_ref(value, slide_count, field="slide"):
    if isinstance(value, str) and value.strip().lower() == "all":
        return "all"
    if isinstance(value, bool) or not isinstance(value, int):
        raise ValueError("%s must be a 1-based slide number or 'all'" % field)
    if value < 1 or value > slide_count:
        raise ValueError("%s %s is out of range (1..%s)" % (field, value, slide_count))
    return value


def validate_ops(ops):
    if not isinstance(ops, list) or not ops:
        raise ValueError("ops must be a non-empty list")
    if len(ops) > MAX_OPS:
        raise ValueError("ops list exceeds limit of %d" % MAX_OPS)
    allowed = {"replace_text", "insert_image", "speaker_notes.set", "slide.create", "slide.delete", "slide.reorder", "shape.create", "shape.text.set", "shape.style.set", "shape.geometry.set"}
    for op in ops:
        if not isinstance(op, dict):
            raise ValueError("each op must be an object")
        name = op.get("op")
        if name not in allowed:
            raise ValueError("unsupported op: %s" % (name,))
        if name == "replace_text":
            if not isinstance(op.get("find"), str) or not op["find"]:
                raise ValueError("replace_text requires a non-empty find string")
            if len(op["find"]) > 500:
                raise ValueError("replace_text find string too long")
            if not isinstance(op.get("replace"), str):
                raise ValueError("replace_text requires a replace string")
            if len(op["replace"]) > 2000:
                raise ValueError("replace_text replace string too long")
            if "slide" not in op:
                raise ValueError("replace_text requires slide")
            if not isinstance(op["slide"], (int, str)) or isinstance(op["slide"], bool):
                raise ValueError("replace_text slide must be a 1-based number or 'all'")
        elif name == "insert_image":
            if "slide" not in op or "asset" not in op:
                raise ValueError("insert_image requires slide and asset")
        elif name == "speaker_notes.set":
            if "slide" not in op or not isinstance(op.get("text"), str):
                raise ValueError("speaker_notes.set requires slide and text")
            if len(op["text"]) > 5000:
                raise ValueError("speaker_notes.set text too long")
        elif name == "slide.create":
            index = op.get("index")
            if index is not None and (not isinstance(index, int) or isinstance(index, bool) or index < 1):
                raise ValueError("slide.create index must be a 1-based integer")
            title = op.get("title")
            if title is not None and (not isinstance(title, str) or len(title) > 300):
                raise ValueError("slide.create title must be a string of at most 300 chars")
        elif name == "slide.delete":
            if "slide" not in op:
                raise ValueError("slide.delete requires slide")
        elif name == "slide.reorder":
            if "slide" not in op or "to" not in op:
                raise ValueError("slide.reorder requires slide and to")
        elif name in {"shape.create", "shape.text.set", "shape.style.set", "shape.geometry.set"}:
            if "slide" not in op or isinstance(op["slide"], bool) or not isinstance(op["slide"], int):
                raise ValueError(name + " requires a 1-based integer slide")
            if name != "shape.create" and not (isinstance(op.get("shape"), (str, int)) and not isinstance(op.get("shape"), bool)):
                raise ValueError(name + " requires a shape name or 1-based index")
            if name == "shape.create" and op.get("kind") not in {"rect", "round_rect", "ellipse", "line", "text"}:
                raise ValueError("shape.create kind must be rect, round_rect, ellipse, line, or text")
            for key in ("left_in", "top_in", "width_in", "height_in", "rotation"):
                if key in op and (isinstance(op[key], bool) or not isinstance(op[key], (int, float)) or not math.isfinite(op[key]) or abs(op[key]) > 10000):
                    raise ValueError(key + " must be a finite number within +/-10000")
            if any(key in op and op[key] <= 0 for key in ("width_in", "height_in")):
                raise ValueError("width_in and height_in must be positive")
            for key in ("fill", "line", "font_color"):
                if key in op and (not isinstance(op[key], str) or not re.fullmatch(r"#?[0-9a-fA-F]{6}", op[key])):
                    raise ValueError(key + " must be a six-digit hex color")
            if name == "shape.text.set" and (not isinstance(op.get("text"), str) or len(op["text"]) > 10000):
                raise ValueError("shape.text.set requires text up to 10000 chars")
            if name == "shape.create" and (not all(isinstance(op.get(k), (int, float)) and not isinstance(op.get(k), bool) and math.isfinite(op[k]) and op[k] > 0 for k in ("width_in", "height_in"))):
                raise ValueError("shape.create requires positive width_in and height_in")
            if name == "shape.create" and ("name" in op and (not isinstance(op["name"], str) or not op["name"].strip() or len(op["name"]) > 100)):
                raise ValueError("shape.create name must be a non-empty string up to 100 chars")
            if name == "shape.create" and ("text" in op and (not isinstance(op["text"], str) or len(op["text"]) > 10000 or op["kind"] == "line")):
                raise ValueError("shape.create text must be up to 10000 chars and cannot be set on a line")
    return ops


def move_slide_id(prs, from_idx0, to_idx0):
    sld_id_lst = prs.slides._sldIdLst
    elements = list(sld_id_lst)
    element = elements.pop(from_idx0)
    sld_id_lst.remove(element)
    sld_id_lst.insert(to_idx0, element)


def delete_slide_at(prs, idx0):
    sld_id_lst = prs.slides._sldIdLst
    sld_id = sld_id_lst[idx0]
    r_id = sld_id.rId
    sld_id_lst.remove(sld_id)
    try:
        prs.part.drop_rel(r_id)
    except Exception:
        pass


def apply_replace_text(prs, op):
    from pptx import Presentation  # noqa: F401  (ensures backend present)
    slide_ref = op["slide"]
    find = op["find"]
    replace = op["replace"]
    targets = list(prs.slides) if (isinstance(slide_ref, str) and slide_ref.lower() == "all") else [prs.slides[slide_ref - 1]]
    count = 0
    for slide in targets:
        for shape in slide.shapes:
            try:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        for run in paragraph.runs:
                            if find in run.text:
                                occurrences = run.text.count(find)
                                run.text = run.text.replace(find, replace)
                                count += occurrences
                if shape.has_table:
                    for row in shape.table.rows:
                        for cell in row.cells:
                            for paragraph in cell.text_frame.paragraphs:
                                for run in paragraph.runs:
                                    if find in run.text:
                                        occurrences = run.text.count(find)
                                        run.text = run.text.replace(find, replace)
                                        count += occurrences
            except Exception:
                continue
    return count


def apply_insert_image(prs, op):
    from pptx.util import Inches

    slide_num = op["slide"]
    if isinstance(slide_num, str) or isinstance(slide_num, bool) or not isinstance(slide_num, int):
        raise ValueError("insert_image slide must be a 1-based slide number")
    if slide_num < 1 or slide_num > len(prs.slides):
        raise ValueError("insert_image slide %s out of range (1..%s)" % (slide_num, len(prs.slides)))
    asset = scoped_resolve(op["asset"], allowed_exts=IMAGE_EXTS, must_exist=True)
    slide = prs.slides[slide_num - 1]

    def inches(value, default):
        try:
            number = float(value)
        except (TypeError, ValueError):
            return Inches(default)
        if number <= 0 or number > 20:
            raise ValueError("image geometry must be within 0..20 inches")
        return Inches(number)

    left = inches(op.get("left_in"), 1.0)
    top = inches(op.get("top_in"), 1.0)
    width = inches(op.get("width_in"), 4.0) if op.get("width_in") is not None else None
    height = inches(op.get("height_in"), 3.0) if op.get("height_in") is not None else None
    if width is not None and height is not None:
        slide.shapes.add_picture(str(asset), left, top, width=width, height=height)
    elif width is not None:
        slide.shapes.add_picture(str(asset), left, top, width=width)
    elif height is not None:
        slide.shapes.add_picture(str(asset), left, top, height=height)
    else:
        slide.shapes.add_picture(str(asset), left, top)
    return {"slide": slide_num, "asset": str(asset)}


def find_shape(slide, selector):
    if isinstance(selector, int) and not isinstance(selector, bool):
        if not 1 <= selector <= len(slide.shapes):
            raise ValueError("shape index is out of range")
        return slide.shapes[selector - 1]
    if isinstance(selector, str) and selector.strip():
        for shape in slide.shapes:
            if shape.name == selector:
                return shape
    raise ValueError("shape name or index was not found")


def apply_shape_op(prs, op):
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt

    slide_num = op["slide"]
    if not 1 <= slide_num <= len(prs.slides):
        raise ValueError("shape operation slide is out of range")
    slide = prs.slides[slide_num - 1]
    name = op["op"]
    if name == "shape.create":
        kind = op["kind"]
        left, top = Inches(float(op.get("left_in", 1))), Inches(float(op.get("top_in", 1)))
        width, height = Inches(float(op["width_in"])), Inches(float(op["height_in"]))
        if kind == "text":
            shape = slide.shapes.add_textbox(left, top, width, height)
        elif kind == "line":
            shape = slide.shapes.add_connector(1, left, top, left + width, top + height)
        else:
            types = {"rect": MSO_SHAPE.RECTANGLE, "round_rect": MSO_SHAPE.ROUNDED_RECTANGLE, "ellipse": MSO_SHAPE.OVAL}
            shape = slide.shapes.add_shape(types[kind], left, top, width, height)
        if op.get("name"):
            shape.name = op["name"]
        if "text" in op:
            shape.text_frame.text = op["text"]
    else:
        shape = find_shape(slide, op["shape"])
    if name in {"shape.create", "shape.geometry.set"}:
        for key, attr in (("left_in", "left"), ("top_in", "top"), ("width_in", "width"), ("height_in", "height")):
            if key in op:
                setattr(shape, attr, Inches(float(op[key])))
        if "rotation" in op:
            shape.rotation = float(op["rotation"])
    if name in {"shape.create", "shape.style.set"}:
        for key, target in (("fill", "fill"), ("line", "line")):
            if key not in op:
                continue
            color = RGBColor.from_string(op[key].lstrip("#").upper())
            if target == "fill":
                shape.fill.solid()
                shape.fill.fore_color.rgb = color
            else:
                shape.line.color.rgb = color
        if "font_color" in op or "font_size" in op:
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if "font_color" in op:
                        run.font.color.rgb = RGBColor.from_string(op["font_color"].lstrip("#").upper())
                    if "font_size" in op:
                        size = op["font_size"]
                        if isinstance(size, bool) or not isinstance(size, (int, float)) or not 1 <= size <= 400:
                            raise ValueError("font_size must be between 1 and 400 points")
                        run.font.size = Pt(size)
    if name == "shape.text.set":
        if not shape.has_text_frame:
            raise ValueError("selected shape does not contain text")
        shape.text_frame.text = op["text"]
    return {"slide": slide_num, "shape": shape.name, "op": name}


def apply_notes_set(prs, op):
    slide_num = op["slide"]
    if isinstance(slide_num, bool) or not isinstance(slide_num, int):
        raise ValueError("speaker_notes.set slide must be a 1-based slide number")
    if slide_num < 1 or slide_num > len(prs.slides):
        raise ValueError("speaker_notes.set slide %s out of range (1..%s)" % (slide_num, len(prs.slides)))
    slide = prs.slides[slide_num - 1]
    slide.notes_slide.notes_text_frame.text = op["text"]
    return {"slide": slide_num, "characters": len(op["text"])}


def apply_slide_create(prs, op):
    layouts = prs.slide_layouts
    layout = layouts[6] if len(layouts) > 6 else layouts[-1]
    new_slide = prs.slides.add_slide(layout)
    new_index = len(prs.slides)  # 1-based position after append
    title = op.get("title")
    if isinstance(title, str) and title:
        try:
            if new_slide.shapes.title is not None:
                new_slide.shapes.title.text = title
            else:
                box = new_slide.shapes.add_textbox(0, 0, 9144000, 1000000)
                box.text_frame.text = title
        except Exception:
            pass
    requested = op.get("index")
    if requested is not None:
        if requested < 1 or requested > len(prs.slides):
            raise ValueError("slide.create index %s out of range (1..%s)" % (requested, len(prs.slides)))
        if requested != new_index:
            move_slide_id(prs, new_index - 1, requested - 1)
            new_index = requested
    return {"index": new_index}


def apply_slide_delete(prs, op):
    slide_num = op["slide"]
    if isinstance(slide_num, bool) or not isinstance(slide_num, int):
        raise ValueError("slide.delete slide must be a 1-based slide number")
    if slide_num < 1 or slide_num > len(prs.slides):
        raise ValueError("slide.delete slide %s out of range (1..%s)" % (slide_num, len(prs.slides)))
    if len(prs.slides) <= 1:
        raise ValueError("slide.delete refused: presentation must keep at least one slide")
    delete_slide_at(prs, slide_num - 1)
    return {"deleted": slide_num}


def apply_slide_reorder(prs, op):
    from_idx = op["slide"]
    to_idx = op["to"]
    for label, value in (("slide", from_idx), ("to", to_idx)):
        if isinstance(value, bool) or not isinstance(value, int):
            raise ValueError("slide.reorder %s must be a 1-based slide number" % label)
        if value < 1 or value > len(prs.slides):
            raise ValueError("slide.reorder %s %s out of range (1..%s)" % (label, value, len(prs.slides)))
    if from_idx != to_idx:
        move_slide_id(prs, from_idx - 1, to_idx - 1)
    return {"from": from_idx, "to": to_idx}


def do_read(payload):
    from pptx import Presentation

    raw_path = payload.get("path", payload.get("input_path", ""))
    path = scoped_resolve(raw_path, allowed_exts=READ_EXTS, must_exist=True)
    prs = Presentation(str(path))
    summary = summarize(prs, path)
    return {
        **summary,
        "backend": "python-pptx",
        "macros_executed": False,
        "verified": True,
        "verification": "application_state",
    }


def do_batch_edit(payload):
    from pptx import Presentation

    raw_path = payload.get("path", payload.get("input_path", ""))
    path = scoped_resolve(raw_path, allowed_exts=READ_EXTS, must_exist=True)
    ops = validate_ops(payload.get("ops", []))
    output_raw = payload.get("output_path")
    if output_raw is not None:
        output_path = scoped_resolve(output_raw, allowed_exts=WRITE_EXTS, must_exist=False)
    else:
        output_path = path
        if output_path.suffix.lower() != ".pptx":
            raise ValueError("in-place batch_edit requires a .pptx file; use output_path for .pptm sources")

    sha_before = sha256_of(path)
    size_before = path.stat().st_size
    backup_path = Path(str(path) + BACKUP_SUFFIX)
    shutil.copy2(path, backup_path)

    prs = Presentation(str(path))
    slide_count_before = len(prs.slides)
    # Sequential pre-validation against a simulated slide count so a
    # batch may create slides and then reference them, while a bad
    # reference is still rejected before anything is written.
    simulated = slide_count_before
    for op in ops:
        name = op["op"]
        if name == "replace_text":
            ref = op["slide"]
            if not (isinstance(ref, str) and ref.lower() == "all"):
                parse_slide_ref(ref, simulated)
        elif name in ("insert_image", "speaker_notes.set", "slide.delete"):
            parse_slide_ref(op["slide"], simulated)
        elif name == "slide.reorder":
            parse_slide_ref(op["slide"], simulated)
            parse_slide_ref(op["to"], simulated)
        elif name in {"shape.create", "shape.text.set", "shape.style.set", "shape.geometry.set"}:
            parse_slide_ref(op["slide"], simulated)
        elif name == "slide.create":
            requested = op.get("index")
            if requested is not None:
                if requested < 1 or requested > simulated + 1:
                    raise ValueError("slide.create index %s out of range (1..%s)" % (requested, simulated + 1))
            simulated += 1
        if name == "slide.delete":
            simulated -= 1
            if simulated < 1:
                raise ValueError("slide.delete refused: presentation must keep at least one slide")

    applied = []
    expected_texts = []
    for op in ops:
        name = op["op"]
        if name == "replace_text":
            replaced = apply_replace_text(prs, op)
            applied.append({"op": name, "replaced_occurrences": replaced})
            if op["replace"]:
                expected_texts.append(op["replace"])
        elif name == "insert_image":
            applied.append({"op": name, **apply_insert_image(prs, op)})
        elif name == "speaker_notes.set":
            applied.append({"op": name, **apply_notes_set(prs, op)})
            expected_texts.append(op["text"])
        elif name == "slide.create":
            applied.append({"op": name, **apply_slide_create(prs, op)})
        elif name == "slide.delete":
            applied.append({"op": name, **apply_slide_delete(prs, op)})
        elif name == "slide.reorder":
            applied.append({"op": name, **apply_slide_reorder(prs, op)})
        elif name in {"shape.create", "shape.text.set", "shape.style.set", "shape.geometry.set"}:
            applied.append(apply_shape_op(prs, op))

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_path))

    # Reopen/readback verification: existence, nonzero size, slide count,
    # expected text presence, and before/after hashes.
    if not output_path.is_file():
        raise RuntimeError("batch_edit save produced no file")
    size_after = output_path.stat().st_size
    if size_after == 0:
        raise RuntimeError("batch_edit save produced an empty file")
    sha_after = sha256_of(output_path)
    reprobe = Presentation(str(output_path))
    slide_count_after = len(reprobe.slides)
    combined = []
    for slide in reprobe.slides:
        combined.extend(slide_text_content(slide))
        combined.append(notes_text(slide))
    haystack = "\n".join(combined)
    missing = [text for text in expected_texts if text and text not in haystack]
    expected_shapes = {}
    for op, applied_op in zip(ops, applied):
        if op["op"] not in {"shape.create", "shape.text.set", "shape.style.set", "shape.geometry.set"}:
            continue
        selector = applied_op.get("shape") if op["op"] == "shape.create" else op["shape"]
        key = (op["slide"], selector)
        expected = expected_shapes.setdefault(key, {})
        if op["op"] == "shape.create":
            for prop, default in (("left_in", 1), ("top_in", 1), ("width_in", None), ("height_in", None)):
                expected[prop] = op.get(prop, default)
        if op["op"] in {"shape.create", "shape.geometry.set"}:
            expected.update({field: op[field] for field in ("left_in", "top_in", "width_in", "height_in", "rotation") if field in op})
        if op["op"] in {"shape.create", "shape.style.set"}:
            expected.update({field: op[field] for field in ("fill", "line", "font_color", "font_size") if field in op})
        if op["op"] == "shape.text.set" or (op["op"] == "shape.create" and "text" in op):
            if "text" in op:
                expected["text"] = op["text"]
    shape_mismatches = []
    for (slide_num, selector), expected in expected_shapes.items():
        try:
            shape = find_shape(reprobe.slides[slide_num - 1], selector)
            checks = {}
            from pptx.util import Inches, Pt
            for field, attr in (("left_in", "left"), ("top_in", "top"), ("width_in", "width"), ("height_in", "height")):
                if field in expected and expected[field] is not None:
                    checks[field] = getattr(shape, attr) == Inches(float(expected[field]))
            if "rotation" in expected:
                checks["rotation"] = abs(float(shape.rotation) - float(expected["rotation"])) < 0.001
            if "text" in expected:
                checks["text"] = shape.has_text_frame and shape.text_frame.text == expected["text"]
            if "fill" in expected:
                checks["fill"] = str(shape.fill.fore_color.rgb).upper() == expected["fill"].lstrip("#").upper()
            if "line" in expected:
                checks["line"] = str(shape.line.color.rgb).upper() == expected["line"].lstrip("#").upper()
            if "font_color" in expected or "font_size" in expected:
                runs = [run for paragraph in shape.text_frame.paragraphs for run in paragraph.runs]
                if "font_color" in expected:
                    checks["font_color"] = bool(runs) and str(runs[0].font.color.rgb).upper() == expected["font_color"].lstrip("#").upper()
                if "font_size" in expected:
                    checks["font_size"] = bool(runs) and runs[0].font.size == Pt(expected["font_size"])
            shape_mismatches.extend({"slide": slide_num, "shape": str(selector), "field": field} for field, value in checks.items() if not value)
        except Exception as exc:
            shape_mismatches.append({"slide": slide_num, "shape": str(selector), "field": str(exc)})
    deletes = sum(1 for entry in applied if entry.get("op") == "slide.delete")
    creates = sum(1 for entry in applied if entry.get("op") == "slide.create")
    reorders = [entry for entry in applied if entry.get("op") == "slide.reorder"]
    expected_count = slide_count_before + creates - deletes
    verified = (
        slide_count_after == expected_count
        and not missing
        and not shape_mismatches
        and size_after > 0
    )
    return {
        "backend": "python-pptx",
        "input_path": str(path),
        "output_path": str(output_path),
        "backup_path": str(backup_path),
        "slide_count_before": slide_count_before,
        "slide_count_after": slide_count_after,
        "slide_count_expected": expected_count,
        "reorder_applied": reorders,
        "applied": applied,
        "missing_texts": missing,
        "shape_mismatches": shape_mismatches,
        "sha256_before": sha_before,
        "sha256_after": sha_after,
        "size_before": size_before,
        "size_after": size_after,
        "macros_executed": False,
        "verified": verified,
        "verification": "persisted_artifact",
    }


def do_export(payload):
    raw_path = payload.get("path", payload.get("input_path", ""))
    path = scoped_resolve(raw_path, allowed_exts=READ_EXTS, must_exist=True)
    fmt = str(payload.get("format", "")).strip().lower()
    if fmt not in EXPORT_FORMATS:
        raise ValueError("export format must be one of ['pdf', 'pptx']")
    output_raw = payload.get("output_path", "")
    if fmt == "pdf":
        output_path = scoped_resolve(output_raw, allowed_exts={".pdf"}, must_exist=False)
        converter = soffice_bin()
        if converter is None:
            raise RuntimeError("pdf_converter_unavailable: no soffice/libreoffice converter found (set COMPTROL_SOFFICE_BIN)")
        output_path.parent.mkdir(parents=True, exist_ok=True)
        completed = subprocess.run(
            [converter, "--headless", "--convert-to", "pdf", "--outdir", str(output_path.parent), str(path)],
            capture_output=True,
            text=True,
            timeout=90,
            check=False,
        )
        if completed.returncode != 0:
            raise RuntimeError("pdf conversion failed: %s" % completed.stderr[-2000:])
        produced = output_path.parent / (path.stem + ".pdf")
        if not produced.is_file():
            raise RuntimeError("pdf conversion produced no file")
        if produced.resolve() != output_path.resolve():
            shutil.move(str(produced), str(output_path))
        size = output_path.stat().st_size if output_path.is_file() else 0
        verified = output_path.is_file() and size > 0
        return {
            "backend": "soffice-pdf",
            "input_path": str(path),
            "output_path": str(output_path),
            "output_size": size,
            "sha256": sha256_of(output_path) if verified else "",
            "converter": converter,
            "verified": verified,
            "verification": "persisted_artifact",
        }
    # pptx export: reopen through python-pptx and save a clean copy so the
    # artifact is proven readable, not just byte-copied.
    from pptx import Presentation

    output_path = scoped_resolve(output_raw, allowed_exts=WRITE_EXTS, must_exist=False)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs = Presentation(str(path))
    prs.save(str(output_path))
    size = output_path.stat().st_size if output_path.is_file() else 0
    verified = output_path.is_file() and size > 0
    reprobe = Presentation(str(output_path)) if verified else None
    return {
        "backend": "python-pptx",
        "input_path": str(path),
        "output_path": str(output_path),
        "output_size": size,
        "sha256": sha256_of(output_path) if verified else "",
        "slide_count": len(reprobe.slides) if reprobe is not None else 0,
        "macros_executed": False,
        "verified": verified,
        "verification": "persisted_artifact",
    }


def handler(request):
    method = request.get("method")
    pptx_ok = has_pptx()
    converter = soffice_bin()
    if method == "handshake":
        backend = "python-pptx" if pptx_ok else "unavailable"
        modes = ["offline"] if pptx_ok else []
        health = "available" if pptx_ok else "unsupported"
        return response(
            request,
            True,
            health,
            {
                "adapter": ADAPTER_ID,
                "backend": backend,
                "modes": modes,
                "route": "offline_open_xml",
                "pptx_available": pptx_ok,
                "pdf_converter": converter or "",
                "pdf_available": converter is not None,
                "macros_executed": False,
            },
        )
    if method == "capabilities":
        return response(
            request,
            True,
            "available" if pptx_ok else "unsupported",
            {
                "backend": "python-pptx" if pptx_ok else "unavailable",
                "pptx_available": pptx_ok,
                "pdf_converter": converter or "",
                "intents": ["presentation.read", "presentation.batch_edit", "presentation.export"],
            },
        )
    if method == "shutdown":
        return response(request, True, "available", {"stopped": True})
    payload = request.get("payload", {}) if isinstance(request.get("payload"), dict) else {}
    intent = payload.get("intent")
    try:
        reject_macro_requests(payload)
        if not pptx_ok:
            return response(
                request,
                False,
                "unsupported",
                error={"code": "pptx_unavailable", "message": "python-pptx is not installed; offline Open XML route unavailable"},
            )
        if intent == "presentation.read":
            result = do_read(payload)
            return response(request, True, "available", result)
        if intent == "presentation.batch_edit":
            result = do_batch_edit(payload)
            health = "available" if result.get("verified") else "degraded"
            return response(request, bool(result.get("verified")), health, result)
        if intent == "presentation.export":
            result = do_export(payload)
            health = "available" if result.get("verified") else "degraded"
            return response(request, bool(result.get("verified")), health, result)
        return response(request, False, "unsupported", error={"code": "unsupported_intent", "message": str(intent)})
    except ImportError as exc:
        return response(request, False, "unsupported", error={"code": "pptx_unavailable", "message": str(exc)})
    except ValueError as exc:
        message = str(exc)
        code = "macro_execution_refused" if message.startswith("macro_execution_refused") else "invalid_request"
        if message.startswith("path_out_of_scope"):
            code = "path_out_of_scope"
        return response(request, False, "unhealthy", error={"code": code, "message": message})
    except RuntimeError as exc:
        message = str(exc)
        if message.startswith("pdf_converter_unavailable"):
            return response(request, False, "unsupported", error={"code": "pdf_converter_unavailable", "message": message})
        return response(request, False, "degraded", error={"code": "presentation_request_failed", "message": message})
    except Exception as exc:
        return response(request, False, "unhealthy", error={"code": "presentation_request_failed", "message": str(exc)})


serve(handler)
